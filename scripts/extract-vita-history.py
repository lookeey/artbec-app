"""Extract only the Vita April-August 2026 delivery history; never edit the XLSX."""
import argparse
import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

SHEETS = [('V. Abr Vita ', '2026-04'), ('V. Maio Vita', '2026-05'),
          ('V. Junho Vita', '2026-06'), ('V. Julho Vita', '2026-07'), ('V. Agosto Vita', '2026-08')]
EXPECTED = {'2026-04': 10, '2026-05': 10, '2026-06': 6, '2026-07': 13, '2026-08': 11}
WEEKDAYS = ['segunda', 'terça', 'quarta', 'quinta', 'sexta', 'sábado', 'domingo']

def extract(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sales = []
    try:
        for sheet_name, period in SHEETS:
            sheet = workbook[sheet_name]
            # These are the two delivery tables, not the monthly quantity summaries.
            for start, end, client in [(13, 24, 'Vita Morumbi'), (28, 42, 'Vita Higienópolis')]:
                for row in sheet.iter_rows(min_row=start, max_row=end, min_col=5, max_col=7):
                    delivery, description, amount = (cell.value for cell in row)
                    if not description:
                        continue
                    if not isinstance(delivery, (date, datetime)):
                        raise ValueError(f'Data ausente/inválida: {sheet_name}!E{row[0].row}')
                    iso = delivery.strftime('%Y-%m-%d')
                    if not iso.startswith(period):
                        raise ValueError(f'Data fora da competência: {sheet_name}!E{row[0].row}')
                    slug = re.sub(r'[^a-z0-9]+', '-', unicodedata.normalize('NFKD', sheet_name.strip()).encode('ascii', 'ignore').decode().lower()).strip('-')
                    source_row = row[0].row
                    sales.append(dict(id=f'sale-{slug}-{source_row}', date=iso, dateNote='', period=period,
                        channel='Vita', client=client, product=str(description).strip(), quantity=None, unitPrice=None,
                        total=None if amount is None else int(Decimal(str(amount))*100), status='review',
                        deliveryStatus='delivered', deliveryDate=iso, deliveredAt='', deliveryWeekday=WEEKDAYS[delivery.weekday()],
                        stockItems=[], payment='', billingClient='', billingDueDay=None,
                        source=sheet_name.strip(), sourceRow=source_row, version=1))
        counts = {period: sum(sale['period'] == period for sale in sales) for _, period in SHEETS}
        if counts != EXPECTED:
            raise ValueError(f'Contagens diferentes das abas conferidas: {counts}')
        return sales
    finally:
        workbook.close()

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--compare-workbook', type=Path)
    args = parser.parse_args()
    sales = extract(args.workbook)
    if args.compare_workbook and sales != extract(args.compare_workbook):
        raise ValueError('As versões da planilha divergem nas saídas de abril a agosto.')
    payload = dict(format='artbec-vita-history-v1', source=dict(file=args.workbook.name,
        sha256=hashlib.sha256(args.workbook.read_bytes()).hexdigest(), sheets=[name for name, _ in SHEETS]), sales=sales)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2)+'\n', encoding='utf-8', newline='\n')
    print(json.dumps(dict(records=len(sales), months=EXPECTED, missingAmounts=sum(s['total'] is None for s in sales)), ensure_ascii=False))
