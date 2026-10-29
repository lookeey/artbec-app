"""Match September 2026 sheet rows to old fuel tombstones without changing XLSX."""
import argparse
import hashlib
import io
import json
import unicodedata
from datetime import date,datetime
from decimal import Decimal,ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('workbook',type=Path)
parser.add_argument('output',type=Path)
args=parser.parse_args()
root=Path(__file__).resolve().parent.parent
base=json.loads((root/'migration/initial/base/operations.json').read_text(encoding='utf-8'))
dump=json.loads((root/'migration/initial/database/database.json').read_text(encoding='utf-8'))
owners={'Esdras','Genilson','Kevin','Nicolas','Vasco'}
def normalize(value):return ' '.join(''.join(c for c in unicodedata.normalize('NFD',value.strip().lower()) if not unicodedata.combining(c)).split())
def key(item):return (normalize(item['technician']),item['date'],item['plate'].strip().upper(),item['amount'])
raw=args.workbook.read_bytes();w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
source='Controle Gasolina.2026';sheet=w[source];rows=[]
for cells in sheet.iter_rows(min_row=11,max_row=sheet.max_row,min_col=5,max_col=8):
 plate,owner,amount,when=(c.value for c in cells)
 if owner not in owners or not isinstance(when,(date,datetime)) or not when.strftime('%Y-%m-%d').startswith('2026-09') or not isinstance(amount,(int,float)):continue
 rows.append(dict(technician=owner,date=when.strftime('%Y-%m-%d'),plate=str(plate or '').strip().upper(),amount=int((Decimal(str(amount))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP)),worksheetRow=cells[0].row))
w.close()
items={item['id']:item for item in base['fuel'] if item['source']==source and item['date'].startswith('2026-09') and item['technician'] in owners}
targets=[]
for deleted in dump['tables']['operation_changes']['rows']:
 if deleted['entity']!='fuel' or not deleted['deleted'] or deleted['id'] not in items:continue
 item=items[deleted['id']]
 matches=[row for row in rows if key(row)==key(item)]
 if len(matches)!=1:raise ValueError(f'Sem correspondência única na planilha: {item["id"]}')
 targets.append(dict(record=item,worksheetRow=matches[0]['worksheetRow'],expectedDeleted=dict(version=deleted['version'],updatedAt=deleted['updated_at'],data=deleted['data'])))
targets.sort(key=lambda target:target['record']['id'])
if len(targets)!=14:raise ValueError(f'Esperadas 14 exclusões antigas, encontradas {len(targets)}')
if len(rows)!=26:raise ValueError(f'Esperados 26 abastecimentos dos responsáveis cadastrados, encontrados {len(rows)}')
summary={owner:dict(count=sum(row['technician']==owner for row in rows),amount=sum(row['amount'] for row in rows if row['technician']==owner)) for owner in sorted(owners)}
payload=dict(format='artbec-fuel-september-restore-v1',source=dict(file=args.workbook.name,sha256=hashlib.sha256(raw).hexdigest(),sheet=source,period='2026-09'),expectedCurrent=summary,targets=targets)
args.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps(dict(oldDeletions=len(targets),worksheetCurrent=summary),ensure_ascii=False))
