"""Extract explicit Well oil-change annotations, preserving the source workbook."""
import argparse,hashlib,io,json,re,unicodedata
from datetime import date,datetime
from pathlib import Path
from openpyxl import load_workbook
p=argparse.ArgumentParser(description=__doc__);p.add_argument('workbook',type=Path);p.add_argument('output',type=Path);a=p.parse_args()
root=Path(__file__).resolve().parent.parent
base=json.loads((root/'migration/initial/base/operations.json').read_text(encoding='utf-8'))
dump=json.loads((root/'migration/initial/database/database.json').read_text(encoding='utf-8'))
recovery=json.loads((root/'migration/oil-recovery-2026-10-29.json').read_text(encoding='utf-8'))
active={r['id']:dict(record=r,expected=None) for r in base['oil']}
for row in dump['tables']['operation_changes']['rows']:
 if row['entity']!='oil':continue
 if row['deleted']:active.pop(row['id'],None)
 else:active[row['id']]=dict(record={**active.get(row['id'],{}).get('record',{}),**json.loads(row['data']),'version':row['version']},expected=row)
for t in recovery['targets']:
 record=t['record'];active[record['id']]=dict(record={**record,'version':t['expectedDeleted']['version']+1} if 'expectedDeleted' in t else record,expected='oil-recovery')
def name(s):
 k=''.join(c for c in unicodedata.normalize('NFD',str(s or '').lower()) if not unicodedata.combining(c)).strip(' ]')
 return {'john':'johnn'}.get(k,k)
def plate(s):return re.sub('[^A-Z0-9]','',str(s or '').upper()).replace('FUAE40','FUA2E40')
raw=a.workbook.read_bytes();w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);s=w['Controle Óleo 26'];targets=[]
for cells in s.iter_rows(min_row=11,min_col=5,max_col=11):
 when,owner,bike,amount,km,last,notes=[c.value for c in cells]
 evidence=[dict(cell=c.coordinate,value=c.value) for c in cells[3:] if isinstance(c.value,str) and re.search(r'\bwell(?:y)?\b',c.value,re.I)]
 if not evidence or not isinstance(when,(date,datetime)) or when.year!=2026:continue
 day=when.strftime('%Y-%m-%d')
 matches=[v for v in active.values() if v['record']['date']==day and name(v['record']['technician'])==name(owner) and plate(v['record']['plate'])==plate(bike) and v['record']['odometer']==km]
 if len(matches)!=1:raise ValueError(f'Conferir registro de óleo da linha {cells[0].row}: {len(matches)} correspondências')
 entry=matches[0];record=entry['record'];old=record['notes']
 # Retain separate historical observations, while replacing shorthand fragments.
 parts=[v.strip() for v in old.split(' · ') if v.strip()]
 retained=[v for v in parts if not re.search(r'\bwell(?:y)?\b|\bwellington\b',v,re.I)]
 updated={**record,'notes':' · '.join(['Trocado no Wellington',*retained]),'entryType':'oil','version':record['version']+1}
 targets.append(dict(before=record,expected=entry['expected'],updated=updated,worksheetRow=cells[0].row,evidence=evidence))
w.close()
if len(targets)!=12:raise ValueError(f'Esperadas 12 indicações explícitas de troca no Well, encontradas {len(targets)}')
payload=dict(format='artbec-oil-wellington-v1',source=dict(file=a.workbook.name,sha256=hashlib.sha256(raw).hexdigest(),sheet='Controle Óleo 26'),targets=targets)
a.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps(dict(targets=len(targets),rows=[t['worksheetRow'] for t in targets]),ensure_ascii=False))
