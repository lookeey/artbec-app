"""Read the updated fuel worksheet and extract additions by content, never by row ID."""
import argparse
import hashlib
import io
import json
import unicodedata
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('workbook',type=Path)
parser.add_argument('output',type=Path)
parser.add_argument('--through',required=True)
args=parser.parse_args()
cutoff=date.fromisoformat(args.through)
root=Path(__file__).resolve().parent.parent
base=json.loads((root/'migration/initial/base/operations.json').read_text(encoding='utf-8'))
sheet_name='Controle Gasolina.2026'
def name_key(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD',value.strip().lower()) if not unicodedata.combining(c)).split())
def key(item):return json.dumps([name_key(item['technician']),item['date'],item['plate'].strip().upper(),item['amount']],ensure_ascii=False,separators=(',',':'))
counts=Counter(key(item) for item in base['fuel'] if item['source']==sheet_name)
aliases={name_key(item['name']):item['name'] for item in base['expenseTechnicians']}
raw=args.workbook.read_bytes();workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
sheet=workbook[sheet_name];candidates=[];excluded=[];valid=0;matched=0;occurrences=Counter()
for cells in sheet.iter_rows(min_row=11,max_row=sheet.max_row,min_col=5,max_col=8):
    plate,owner,amount,when=(cell.value for cell in cells)
    if not any(value is not None for value in [plate,owner,amount,when]):continue
    row=cells[0].row
    if not isinstance(when,(date,datetime)) or not owner or not isinstance(amount,(int,float)):
        excluded.append(dict(sourceRow=row,range=f'E{row}:H{row}',technician=str(owner or ''),rawDate=str(when or ''),reason='Linha sem data ou valor válidos; não corrigida automaticamente'))
        continue
    iso=when.strftime('%Y-%m-%d')
    item=dict(technician=aliases.get(name_key(str(owner)),str(owner).strip()),date=iso,plate=str(plate or '').strip().upper(),amount=int((Decimal(str(amount))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP)))
    fingerprint=key(item);occurrences[fingerprint]+=1;valid+=1
    if counts[fingerprint]:counts[fingerprint]-=1;matched+=1;continue
    if iso<'2026-01-01' or iso>args.through:
        excluded.append(dict(sourceRow=row,range=f'E{row}:H{row}',rawDate=iso,reason='Fora de 2026 ou posterior à data de corte'))
        continue
    identifier='fuel-import-20261029-'+hashlib.sha256((fingerprint+'|'+str(occurrences[fingerprint])).encode('utf-8')).hexdigest()[:24]
    candidates.append(dict(id=identifier,**item,source=sheet_name,sourceRow=row,version=1))
workbook.close()
if any(counts.values()):raise ValueError('A planilha diverge dos abastecimentos antigos. Revise antes de importar.')
if len(candidates)!=38:raise ValueError(f'Foram encontrados {len(candidates)} registros novos; confira a versão da planilha.')
payload=dict(format='artbec-fuel-sheet-recovery-v1',source=dict(file=args.workbook.name,sha256=hashlib.sha256(raw).hexdigest(),sheet=sheet_name,through=args.through),compared=dict(existingRows=matched,validRows=valid),excluded=excluded,fuel=candidates)
args.output.parent.mkdir(parents=True,exist_ok=True)
args.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps(dict(candidates=len(candidates),matched=matched,excluded=excluded,months=dict(Counter(item['date'][:7] for item in candidates))),ensure_ascii=False))
