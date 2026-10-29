"""Read 2026 oil history and approximate change references without editing XLSX."""
import argparse,hashlib,io,json,re,unicodedata,uuid
from datetime import date,datetime
from decimal import Decimal,ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
p=argparse.ArgumentParser(description=__doc__);p.add_argument('workbook',type=Path);p.add_argument('output',type=Path);a=p.parse_args()
root=Path(__file__).resolve().parent.parent
base=json.loads((root/'migration/initial/base/operations.json').read_text(encoding='utf-8'))
dump=json.loads((root/'migration/initial/database/database.json').read_text(encoding='utf-8'))
saved=[r for r in dump['tables']['operation_changes']['rows'] if r['entity'] in ('oil','oilMonitor')]
def name(s):
 s=' '.join(str(s or '').strip(' ]').split());k=''.join(c for c in unicodedata.normalize('NFD',s.lower()) if not unicodedata.combining(c))
 return {'john':'Johnn','simao':'Simão','rogerio':'Rogério','oficina':'Oficina'}.get(k,s.title())
def plate(s):return re.sub('[^A-Z0-9]','',str(s or '').upper()).replace('FUAE40','FUA2E40')
def day(r):return name(r['technician']),r['date'],plate(r['plate'])
def iso(v):return v.strftime('%Y-%m-%d') if isinstance(v,(date,datetime)) else ''
def integer(v):return int(v) if isinstance(v,(float,int)) and v>=0 and float(v).is_integer() else None
raw=a.workbook.read_bytes();w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);sheet=w['Controle Óleo 26'];records=[];excluded=[]
for cells in sheet.iter_rows(min_row=11,min_col=5,max_col=11):
 when,owner,bike,amount,km,last,notes=[c.value for c in cells];d=iso(when);row=cells[0].row
 if not d.startswith('2026-') or d>'2026-10-29':continue
 if not owner or not bike or all(v is None for v in (amount,km,last,notes)):
  excluded.append({'row':row,'reason':'no-reading-change-or-expense'});continue
 note=[str(v).strip() for v in (amount if isinstance(amount,str) and amount.strip()!='-' else None,last if isinstance(last,str) and last.strip()!='-' else None,notes) if v is not None and str(v).strip()]
 item=dict(date=d,technician=name(owner),plate=plate(bike),amount=int((Decimal(str(amount))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP)) if isinstance(amount,(int,float)) else None,odometer=integer(km),lastChange=iso(last),notes=' · '.join(dict.fromkeys(note)),source=sheet.title,sourceRow=row,version=1)
 # A missing price is not a zero expense or proof that the oil was not changed.
 item['entryType']='oil' if item['amount'] is not None or (isinstance(amount,str) and any(x in amount.lower() for x in ('well','troc'))) or bool(item['lastChange']) else 'reading'
 digest=hashlib.sha256((item['plate']+':'+item['date']).encode()).hexdigest()[:32]
 item['id']='oil-'+'-'.join([digest[:8],digest[8:12],digest[12:16],digest[16:20],digest[20:]])
 records.append(item)
w.close()
active={r['id']:r for r in base['oil']};deleted=[]
for row in saved:
 if row['entity']!='oil':continue
 data=json.loads(row['data']);original=active.get(row['id'])
 if row['deleted']:active.pop(row['id'],None);deleted.append((row,original or data))
 else:active[row['id']]={**(original or {}),**data}
keys={day(r) for r in active.values()};targets=[]
for item in records:
 if day(item) in keys:continue
 old=[(row,original) for row,original in deleted if original and day(original)==day(item)]
 original=next(((row,record) for row,record in old if record.get('source')=='Controle Óleo 26'),None)
 if old and not original:excluded.append({'row':item['sourceRow'],'reason':'manual-deletion-preserved'});continue
 target={'record':item}
 if original:
  row,record=original;target['record']={**item,'id':record['id']};target['expectedDeleted']=row;target['original']=record
 targets.append(target)
# One card per current motorcycle; former employees remain only in the ledger.
owners={'FUA2E40':'Esdras','FCC7G74':'Vasco','GDM8147':'Genilson','GIM5J14':'Kevin','FLO3A78':'Nicolas'};monitors=[]
for bike,owner in owners.items():
 rows=[r for r in records if r['plate']==bike and r['odometer'] is not None]
 latest=max(rows,key=lambda r:(r['date'],r['sourceRow']))
 refs=[r for r in rows if r['lastChange'] and r['lastChange']<='2026-10-29']
 latestDate=max(r['lastChange'] for r in refs);oil=min((r for r in refs if r['lastChange']==latestDate),key=lambda r:(r['date'],r['sourceRow']))
 if latest['odometer']<oil['odometer']:raise ValueError('Regressão nas referências atuais de '+bike)
 m=dict(id='oilMonitor-'+str(uuid.uuid5(uuid.NAMESPACE_URL,'artbec-oil-20261029-'+bike)),version=1,technician=owner,plate=bike,currentKm=latest['odometer'],readingDate=max(latest['date'],oil['lastChange']),oilKm=oil['odometer'],oilDate=oil['lastChange'],serviceKm=None,serviceDate='',history=[],importReference=dict(source=sheet.title,oilRow=oil['sourceRow'],oilReadingDate=oil['date'],currentReadingDate=latest['date'],approximate=True,oilKm=oil['odometer'],oilDate=oil['lastChange'],currentKm=latest['odometer'],readingDate=max(latest['date'],oil['lastChange'])))
 for r in sorted(rows,key=lambda r:(r['date'],r['sourceRow'])):
  m['history'].append(dict(date=r['date'],km=r['odometer'],event='reading',sourceRow=r['sourceRow']))
 m['history'].append(dict(date=m['oilDate'],km=m['oilKm'],event='oil',sourceRow=oil['sourceRow'],approximate=True,oilKm=oil['odometer'],oilDate=oil['lastChange'],currentKm=latest['odometer'],readingDate=max(latest['date'],oil['lastChange'])))
 old=next((row for row in saved if row['entity']=='oilMonitor' and not row['deleted'] and plate(json.loads(row['data'])['plate'])==bike),None)
 monitors.append(dict(record=m,expectedExisting=old))
if len(targets)!=23 or sum('expectedDeleted' in t for t in targets)!=4:raise ValueError('Conferir escopo dos lançamentos')
payload=dict(format='artbec-oil-2026-recovery-v1',source=dict(file=a.workbook.name,sha256=hashlib.sha256(raw).hexdigest(),sheet='Controle Óleo 26',through='2026-10-29'),targets=targets,monitors=monitors,excluded=excluded,existingSnapshot=saved)
a.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps(dict(ledgerTargets=len(targets),restorations=sum('expectedDeleted' in t for t in targets),newAmount=sum(t['record']['amount'] or 0 for t in targets if 'expectedDeleted' not in t),monitors=[{k:r['record'][k] for k in ('technician','plate','currentKm','readingDate','oilKm','oilDate')} for r in monitors],excluded=excluded),ensure_ascii=False))
